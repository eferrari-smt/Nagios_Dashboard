"""Migración única: convierte el Excel multi-hoja de inventario a un CSV plano.

El Excel nunca se toca ni se borra — queda como respaldo. Se recorren
todas las hojas (GCBA, Rosario, RESTO, STOCK) en orden y se concatenan
sus filas en un solo archivo, ubicando las columnas por encabezado (no
por posición) porque no todas las hojas tienen las mismas columnas
opcionales (ej. Rosario y STOCK no tienen "Notas").

Uso:
    python migrate_excel_to_csv.py <ruta_excel> <ruta_csv_salida>
"""

from __future__ import annotations

import csv
import sys

import openpyxl

_ALL_COLUMNS = ["Nombre PC", "RAK", "ID AnyDesk", "Cliente", "ID Chrome Remote Desktop", "Notas"]


def migrate(xlsx_path: str, csv_path: str) -> int:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    rows_out = [_ALL_COLUMNS]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        column_index = {name: header.index(name) for name in _ALL_COLUMNS if name in header}

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
                continue

            def cell(name: str) -> str:
                idx = column_index.get(name)
                if idx is None or idx >= len(row) or row[idx] is None:
                    return ""
                return str(row[idx]).strip()

            rows_out.append([cell(name) for name in _ALL_COLUMNS])

    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows_out)

    return len(rows_out) - 1


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Uso: python migrate_excel_to_csv.py <ruta_excel> <ruta_csv_salida>")
        sys.exit(1)

    count = migrate(sys.argv[1], sys.argv[2])
    print(f"Migradas {count} filas a {sys.argv[2]}")
